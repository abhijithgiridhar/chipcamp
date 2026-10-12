#!/usr/bin/env python3
"""Makes one PowerPoint file per deck from session-plan/decks/*.js (the same content as the web slides).
Run from the camp-blocks folder:
  python3 tools/render_art.py     (once, makes the character pictures)
  python3 tools/make_pptx.py
  node build.js"""
import json, math, os, re, subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'session-plan', 'pptx')
ART = os.path.join(OUT, 'art')
W, H = 13.333, 7.5
HEAD, BODY, MONO = 'Century Gothic', 'Calibri', 'Courier New'
PAPER, CARD, LINE, INK, INDIGO, MUSTARD, GREY = 'F0EFE9', 'FAF8F3', 'E3E0D6', '2C2C34', '5B5B9F', 'FDB515', '5A5A63'
EMOJI = {'chipbot': '🤖', 'peeko': '👀', 'jarvis': '🌱'}


def load_decks():
    js = "const o=['kickoff','chipbot','peeko','jarvis'].map(i=>require('./session-plan/decks/'+i+'.js'));process.stdout.write(JSON.stringify(o))"
    return json.loads(subprocess.check_output(['node', '-e', js], cwd=ROOT))


def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def tint(h, t):
    h = h.lstrip('#')
    c = [int(h[i:i + 2], 16) for i in (0, 2, 4)]
    return '%02X%02X%02X' % tuple(round(v + (255 - v) * t) for v in c)


def lines_needed(text, w_in, size, factor):
    per = max(1, int(w_in * 72 / (size * factor)))
    return sum(max(1, math.ceil(len(p) / per)) for p in text.replace('**', '').split('\n'))


def fit(text, w_in, h_in, start, minimum=14, factor=0.62, line=1.2):
    """Biggest font size (pt) at which the text should fit the box. A rough estimate, kept on the safe side."""
    for size in range(start, minimum - 1, -1):
        if lines_needed(text, w_in, size, factor) * size * line / 72 <= h_in:
            return size
    return minimum


def add_runs(par, text, size, color, bold=False, accent=None, font=BODY):
    for p in re.split(r'(\*\*[^*]+\*\*)', text):
        if not p:
            continue
        b = p.startswith('**') and p.endswith('**')
        r = par.add_run()
        r.text = (p[2:-2] if b else p).replace('\u2011', '-')
        r.font.size = Pt(size)
        r.font.name = font
        r.font.bold = bold or b
        r.font.color.rgb = rgb(accent if (b and accent) else color)


def text(slide, s, x, y, w, h, size, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, accent=None, font=BODY, name=None, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(s.split('\n')):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        add_runs(par, line, size, color, bold, accent, font)
    if spacing:
        for r in tf.paragraphs[0].runs:
            r.font._element.set('spc', str(spacing))
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16, line_w=2):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def pic(slide, key, x, y, w, h=None, name=None):
    path = os.path.join(ART, key + '.png')
    if not os.path.exists(path):
        return None
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h) if h else None)
    p.name = name or ('Art ' + key)
    return p


def art(slide, key, cx, cy, d):
    """A character picture centred on (cx, cy), d inches square. Emoji if there is no picture."""
    if os.path.exists(os.path.join(ART, key + '.png')):
        pic(slide, key, cx - d / 2, cy - d / 2, d, d)
    else:
        text(slide, EMOJI.get(key, '⭐'), cx - d / 2, cy - d / 2, d, d, int(d * 40), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name='Art')


def eyebrow(slide, s, deck, idx, color=INDIGO):
    label = s.get('eyebrow')
    if not label:
        acts = [a for a in deck['acts'] if idx >= a['from']]
        label = acts[-1]['name'] if acts else deck['title']
    text(slide, label.upper(), 1.0, 0.62, 9, 0.35, 14, color, bold=True, name='Eyebrow', spacing=300)


def heading(slide, s, color=INK, y=0.98, size=40):
    sz = fit(s, W - 2.0, 1.0, size, 28, factor=0.66)
    text(slide, s, 1.0, y, W - 2.0, 1.0, sz, color, bold=True, anchor=MSO_ANCHOR.MIDDLE, font=HEAD, name='Title')


def pill(slide, label, x, y, w, h, fill=MUSTARD, color=INK, size=16, href=None):
    b = box(slide, x, y, w, h, fill, radius=0.5)
    t = text(slide, label, x, y, w, h, size, color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
    if href and href.startswith('http'):
        b.click_action.hyperlink.address = href
        t.click_action.hyperlink.address = href


def quip(slide, s):
    if not s.get('quip'):
        return
    pic(slide, 'bug', 0.35, 6.4, 0.85, 0.85, name='Bug')
    bw = min(6.2, max(2.6, len(s['quip']) * 0.115 + 0.5))
    box(slide, 1.3, 6.5, bw, 0.55, 'FFFFFF', line=INK, radius=0.4, line_w=2)
    text(slide, s['quip'], 1.45, 6.5, bw - 0.3, 0.55, 14, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE, name='Bug says')


def number(slide, n, color=GREY):
    text(slide, str(n), W - 1.0, 6.95, 0.6, 0.3, 14, color, align=PP_ALIGN.RIGHT, name='Slide number')


def build(deck):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    c = deck['color'].lstrip('#')
    for idx, s in enumerate(deck['slides']):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        t = s['t']
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(MUSTARD if t == 'big' else PAPER)

        if t == 'title':
            pic(sl, 'splash', 5.4, 3.4, 8.6, 4.84, name='Splash')
            text(sl, 'MICRON CHIP CAMP', W - 3.6, 0.5, 2.9, 0.3, 12, INDIGO, bold=True, align=PP_ALIGN.RIGHT, font=HEAD, spacing=200, name='Badge')
            text(sl, s.get('kicker', '').upper(), 1.0, 1.55, 6.4, 0.35, 14, INDIGO, bold=True, name='Eyebrow', spacing=300)
            tsz = fit(s['title'], 6.3, 2.9, 60, 32, factor=0.7)
            text(sl, s['title'], 1.0, 2.0, 6.3, 2.9, tsz, INDIGO, bold=True, font=HEAD, name='Title')
            used = lines_needed(s['title'], 6.3, tsz, 0.7) * tsz * 1.2 / 72
            text(sl, s.get('sub', ''), 1.0, 2.0 + used + 0.25, 6.2, 1.3, fit(s.get('sub', ''), 6.2, 1.3, 24, 16), GREY, name='Subtitle')
            art(sl, s.get('art', 'chip'), 10.3, 3.6, 4.7)
        elif t == 'end':
            pic(sl, 'splash', 5.0, 4.0, 8.8, 4.95, name='Splash')
            eyebrow(sl, s, deck, idx)
            text(sl, s['title'], 1.0, 1.2, 8.8, 2.2, fit(s['title'], 8.8, 2.2, 56, 32, factor=0.7), INDIGO, bold=True, font=HEAD, anchor=MSO_ANCHOR.TOP, name='Title')
            text(sl, s.get('sub', ''), 1.0, 3.3, 7.4, 1.3, fit(s.get('sub', ''), 7.4, 1.3, 24, 16), GREY)
            pic(sl, 'kids', 5.9, 4.1, 7.1, name='The kids')
        elif t == 'big':
            if s.get('eyebrow'):
                eyebrow(sl, s, deck, idx, INK)
            has = bool(s.get('art'))
            w = 8.0 if has else 11.0
            bsz = fit(s['big'], w, 3.4, 66, 36, factor=0.7)
            text(sl, s['big'], 1.0, 1.4, w, 3.4, bsz, INK, bold=True, font=HEAD, anchor=MSO_ANCHOR.BOTTOM, name='Title')
            text(sl, s.get('small', ''), 1.0, 5.0, w, 1.6, fit(s.get('small', ''), w, 1.6, 24, 16), INK)
            if has:
                art(sl, s['art'], 11.0, 4.8, 3.0)
        elif t == 'comic':
            eyebrow(sl, s, deck, idx)
            if s.get('title'):
                heading(sl, s['title'], size=34)
            n = len(s['panels'])
            gap = 0.35
            pw = (W - 2.0 - gap * (n - 1)) / n
            for i, p in enumerate(s['panels']):
                x = 1.0 + i * (pw + gap)
                box(sl, x, 2.1, pw, 4.4, CARD, line=INK, radius=0.07, line_w=2.5)
                b = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, Inches(x + 0.25), Inches(2.35), Inches(pw - 0.5), Inches(1.35))
                b.fill.solid(); b.fill.fore_color.rgb = rgb('FFFFFF'); b.line.color.rgb = rgb(INK); b.line.width = Pt(2.5)
                b.adjustments[0] = 0.0
                b.adjustments[1] = 0.78
                b.shadow.inherit = False
                b.name = 'Speech bubble'
                text(sl, p['say'], x + 0.4, 2.4, pw - 0.8, 1.25, fit(p['say'], pw - 0.8, 1.25, 22, 14, factor=0.6), INK, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                art(sl, p['art'], x + pw / 2, 5.15, 2.0 if n == 3 else 2.3)
                if p.get('name'):
                    text(sl, p['name'].upper(), x, 6.1, pw, 0.3, 12, INDIGO, bold=True, align=PP_ALIGN.CENTER, spacing=150)
        elif t == 'list':
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            n = len(s['items'])
            top, gap = 2.15, 0.2
            ih = min(1.2, (6.2 - top - gap * (n - 1)) / n + 0.0)
            width = 8.7 if n <= 3 else 11.3
            for i, it in enumerate(s['items']):
                y = top + i * (ih + gap)
                box(sl, 1.0, y, width, ih, CARD, line=LINE, radius=0.2)
                lead = it[0]
                box(sl, 1.25, y + (ih - 0.8) / 2, 0.8, 0.8, MUSTARD if (len(lead) <= 3 and lead.isascii()) else tint(c, 0.78), shape=MSO_SHAPE.OVAL)
                text(sl, lead, 1.25, y + (ih - 0.8) / 2, 0.8, 0.8, 26 if len(lead) <= 2 else 11, INK, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                text(sl, it[1], 2.3, y + 0.1, width - 1.6, ih * 0.5, fit(it[1], width - 1.6, ih * 0.5, 22, 15), INDIGO, bold=True, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                text(sl, it[2] or '', 2.3, y + ih * 0.52, width - 1.6, ih * 0.42, fit(it[2] or ' ', width - 1.6, ih * 0.42, 16, 12), GREY)
            if n <= 3 and s.get('art'):
                art(sl, s['art'], 11.2, 4.2, 2.5)
        elif t == 'levels':
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            n = len(s['items']) + 1
            M = 1.2
            step = (W - 2 * M) / (n - 1)
            pts = [(M + step * i, 3.6 if i % 2 == 0 else 4.7) for i in range(n)]
            for i in range(n - 1):
                cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(pts[i][0]), Inches(pts[i][1]), Inches(pts[i + 1][0]), Inches(pts[i + 1][1]))
                cn.line.color.rgb = rgb('B4B4D2'); cn.line.width = Pt(4); cn.line.dash_style = MSO_LINE.ROUND_DOT
            for i, (x, y) in enumerate(pts):
                goal = i == n - 1
                box(sl, x - 0.5, y - 0.5, 1.0, 1.0, INDIGO if goal else MUSTARD, shape=MSO_SHAPE.OVAL)
                if goal:
                    text(sl, '🏆', x - 0.5, y - 0.5, 1.0, 1.0, 30, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                    text(sl, 'Robot alive', x - 1.0, y + 0.6, 2.0, 0.4, 16, INK, bold=True, align=PP_ALIGN.CENTER, font=HEAD)
                else:
                    it = s['items'][i]
                    text(sl, it[0], x - 0.5, y - 0.5, 1.0, 1.0, 30, INK, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                    text(sl, it[1], x - 1.0, y + 0.6, 2.0, 0.4, 16, INK, bold=True, align=PP_ALIGN.CENTER, font=HEAD)
                    text(sl, it[2] or '', x - 1.0, y + 1.0, 2.0, 0.5, 12, GREY, align=PP_ALIGN.CENTER)
        elif t in ('steps', 'activity'):
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            if t == 'activity':
                pill(sl, '⏱ %d min' % s['minutes'], W - 3.3, 0.55, 2.3, 0.75, size=24)
            n = len(s['steps'])
            top, avail = 2.15, 6.5 - 2.15
            rh = min(0.95, avail / n)
            width = 8.7 if n <= 4 else 11.3
            for i, st in enumerate(s['steps']):
                y = top + i * rh
                box(sl, 1.0, y, width, rh - 0.12, CARD, line=LINE, radius=0.25)
                box(sl, 1.15, y + (rh - 0.12 - 0.58) / 2, 0.58, 0.58, MUSTARD, shape=MSO_SHAPE.OVAL)
                text(sl, str(i + 1), 1.15, y + (rh - 0.12 - 0.58) / 2, 0.58, 0.58, 20, INK, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                text(sl, st, 1.95, y, width - 1.15, rh - 0.12, fit(st, width - 1.15, rh - 0.12, 20, 14, factor=0.6), INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            if n <= 4 and s.get('art'):
                art(sl, s['art'], 11.3, 4.3, 2.4)
        elif t == 'parts':
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            cw, ch, gx, gy = 3.55, 2.0, 0.3, 0.25
            for i, p in enumerate(s['parts']):
                x = 1.0 + (i % 3) * (cw + gx)
                y = 2.15 + (i // 3) * (ch + gy)
                box(sl, x, y, cw, ch, CARD, line=LINE, radius=0.14)
                box(sl, x + cw / 2 - 0.4, y + 0.15, 0.8, 0.8, tint(c, 0.8), shape=MSO_SHAPE.OVAL)
                text(sl, p[0], x + cw / 2 - 0.4, y + 0.15, 0.8, 0.8, 26, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, p[1], x + 0.15, y + 1.02, cw - 0.3, 0.4, fit(p[1], cw - 0.3, 0.4, 20, 15), INDIGO, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                text(sl, p[2], x + 0.15, y + 1.42, cw - 0.3, 0.5, fit(p[2], cw - 0.3, 0.5, 14, 11), GREY, align=PP_ALIGN.CENTER)
        elif t == 'demo':
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            text(sl, s['text'], 1.0, 2.3, 5.8, 2.2, fit(s['text'], 5.8, 2.2, 26, 16), INK, bold=True)
            if s.get('link'):
                web = s['link']['href'].startswith('http')
                pill(sl, s['link']['label'] + ' ↗', 1.0, 4.7, 4.6, 0.8, fill=INDIGO, color='FFFFFF', size=18, href=s['link']['href'])
                text(sl, 'Click to open it in your browser.' if web else 'Open this on the camp website. The web version of this slide runs the robot live.', 1.0, 5.7, 5.8, 0.7, 13, GREY)
            box(sl, 7.3, 2.15, 5.0, 4.0, 'FFFFFF', line=INDIGO, radius=0.08, line_w=3)
            key = s.get('art') or (deck.get('stage') if s.get('stage') else 'spark')
            art(sl, key, 9.8, 4.15, 3.4)
        elif t == 'question':
            eyebrow(sl, s, deck, idx)
            text(sl, s['title'], 1.0, 1.0, 8.9, 1.9, fit(s['title'], 8.9, 1.9, 38, 24, factor=0.7), INK, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, name='Title')
            for i, o in enumerate(s['options']):
                y = 3.1 + i * 1.12
                box(sl, 1.0, y, 8.9, 0.95, CARD, line=LINE, radius=0.25, line_w=2.5)
                box(sl, 1.2, y + 0.18, 0.6, 0.6, MUSTARD, shape=MSO_SHAPE.OVAL)
                text(sl, 'ABCD'[i], 1.2, y + 0.18, 0.6, 0.6, 20, INK, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                text(sl, o, 2.05, y, 7.6, 0.95, fit(o, 7.6, 0.95, 22, 14, factor=0.6), INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            if s.get('art'):
                art(sl, s['art'], 11.3, 4.3, 2.6)
        elif t == 'code':
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            nl = s['code'].count('\n') + 1
            ph = min(3.7, 0.55 * nl + 0.5)
            box(sl, 1.0, 2.2, 11.3, ph, INK, radius=0.06)
            text(sl, s['code'], 1.4, 2.4, 10.5, ph - 0.3, 26, 'F2F0FF', font=MONO, name='Code')
            text(sl, s.get('text', ''), 1.0, 2.2 + ph + 0.2, 11.3, 0.5, 20, INK, bold=True)
        elif t == 'crew':
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            for i, cr in enumerate(s['crew']):
                x = 1.0 + i * 3.85
                box(sl, x, 2.15, 3.5, 4.3, CARD, line=LINE, radius=0.12)
                art(sl, cr[0], x + 1.75, 3.5, 2.2)
                text(sl, cr[1], x + 0.15, 4.75, 3.2, 0.6, 26, INDIGO, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                text(sl, cr[2], x + 0.25, 5.4, 3.0, 0.95, fit(cr[2], 3.0, 0.95, 18, 13), GREY, align=PP_ALIGN.CENTER)
        elif t == 'wanted':
            eyebrow(sl, s, deck, idx)
            heading(sl, s['title'])
            pw, gap = 3.55, 0.3
            for i, b in enumerate(s['bugs']):
                x = 1.0 + i * (pw + gap)
                box(sl, x, 2.15, pw, 4.3, 'FBF3DC', line=INK, radius=0.04, line_w=2.5)
                text(sl, 'WANTED', x, 2.25, pw, 0.55, 28, INDIGO, bold=True, align=PP_ALIGN.CENTER, font=HEAD, spacing=250)
                pic(sl, 'bug', x + pw / 2 - 0.65, 2.85, 1.3, 1.3, name='Bug')
                text(sl, b[0], x + 0.15, 4.2, pw - 0.3, 0.55, fit(b[0], pw - 0.3, 0.55, 18, 13, factor=0.66), INK, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
                text(sl, 'Seen: ' + b[1], x + 0.2, 4.8, pw - 0.4, 0.7, fit('Seen: ' + b[1], pw - 0.4, 0.7, 14, 11), GREY, align=PP_ALIGN.CENTER)
                box(sl, x + 0.2, 5.6, pw - 0.4, 0.7, MUSTARD, radius=0.3)
                text(sl, 'Reward: ' + b[2], x + 0.3, 5.6, pw - 0.6, 0.7, fit('Reward: ' + b[2], pw - 0.6, 0.7, 13, 10), INK, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        if t not in ('title', 'end'):
            number(sl, idx + 1, INK if t == 'big' else GREY)
        quip(sl, s)
        if s.get('notes'):
            sl.notes_slide.notes_text_frame.text = s['notes']
    return prs


def main():
    os.makedirs(OUT, exist_ok=True)
    if not os.path.isdir(ART):
        print('No pictures yet: run python3 tools/render_art.py first. Using emoji for now.')
    for deck in load_decks():
        path = os.path.join(OUT, deck['id'] + '.pptx')
        build(deck).save(path)
        print('wrote', os.path.relpath(path, ROOT), '(%d slides)' % len(deck['slides']))


if __name__ == '__main__':
    main()
